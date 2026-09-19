"""数据集管理：列表/预览/用例/上传/删除（P10，删除按 owner-or-admin）。"""
import io
import json
import uuid

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from ...core.database import get_db
from ...core.deps import get_current_user, require_admin, require_owner_or_admin
from ...models import Case, CaseResult, Dataset, User
from ..shared import ok
from .schemas import (
    CaseBatchGetReq,
    CaseListOut,
    CaseOut,
    ColumnMapping,
    DatasetBatchDeleteReq,
    DatasetListOut,
    DatasetNameUpdateReq,
    DatasetOut,
    DatasetPreviewOut,
)

router = APIRouter(prefix="/datasets", tags=["datasets"])

REQUIRED = ["case_no", "user_input", "expected_gt"]


@router.get("")
def list_datasets(
    owner: str | None = None,
    keyword: str | None = None,
    username: str | None = None,
    created_start: str | None = None,
    created_end: str | None = None,
    page: int = 1,
    page_size: int = 20,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    if page < 1:
        page = 1
    if page_size < 1 or page_size > 100:
        page_size = 20
    # 动态拼接 WHERE：owner / 数据集名称 / 归属人 username / 创建时间区间，均为可选
    filters = []
    if owner == "mine":
        filters.append(Dataset.owner_id == _user.id)
    elif owner and owner != "all":
        filters.append(Dataset.owner_id == uuid.UUID(owner))
    if keyword:
        filters.append(Dataset.name.ilike(f"%{keyword}%"))
    if username:
        filters.append(User.username == username)
    if created_start:
        filters.append(Dataset.created_at >= created_start)
    if created_end:
        filters.append(Dataset.created_at <= created_end)
    # 始终 join users 以带出上传人 username（isouter 保证 owner 缺失时仍返回）
    q = db.query(Dataset, User).join(User, Dataset.owner_id == User.id, isouter=True)
    if filters:
        q = q.filter(*filters)
    total = q.count()
    rows = (
        q.order_by(Dataset.created_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    items = [
        DatasetOut.model_validate(d).model_copy(update={"username": u.username})
        for d, u in rows
    ]
    data = DatasetListOut(items=items, page=page, page_size=page_size, total=total)
    return ok(data.model_dump())


@router.get("/uploaders")
def list_uploaders(db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    """返回所有数据集的去重上传人列表（供筛选下拉使用）。"""
    rows = (
        db.query(User.username)
        .join(Dataset, Dataset.owner_id == User.id)
        .distinct()
        .order_by(User.username)
        .all()
    )
    return ok([r[0] for r in rows if r[0]])


@router.get("/{dataset_id}")
def get_dataset(dataset_id: uuid.UUID, db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    d = db.get(Dataset, dataset_id)
    if d is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "数据集不存在")
    return ok(DatasetOut.model_validate(d).model_dump())


@router.get("/{dataset_id}/preview")
def preview(dataset_id: uuid.UUID, db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    d = db.get(Dataset, dataset_id)
    if d is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "数据集不存在")
    rows = db.query(Case).filter(Case.dataset_id == dataset_id).order_by(Case.sort_order).limit(5).all()
    headers = ["用例编号", "测试输入", "预期结果", "评测一级维度", "评测二级维度"]
    rows_data = [[c.case_no, c.user_input, c.expected_gt, c.dimension_l1 or "", c.dimension_l2 or ""] for c in rows]
    data = DatasetPreviewOut(dataset_id=str(d.id), name=d.name, case_count=d.case_count, headers=headers, rows=rows_data)
    return ok(data.model_dump())


@router.get("/{dataset_id}/cases")
def list_cases(
    dataset_id: uuid.UUID,
    page: int = 1,
    page_size: int = 20,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    if page < 1:
        page = 1
    if page_size < 1 or page_size > 100:
        page_size = 20
    q = db.query(Case).filter(Case.dataset_id == dataset_id)
    total = q.count()
    items = q.order_by(Case.sort_order).offset((page - 1) * page_size).limit(page_size).all()
    data = CaseListOut(items=items, page=page, page_size=page_size, total=total)
    return ok(data.model_dump())


@router.post("/{dataset_id}/cases/batch")
def batch_get_cases(
    dataset_id: uuid.UUID,
    req: CaseBatchGetReq,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    """按 case_no 批量查询数据集用例（报告页按需取指定用例的输入/预期）。"""
    d = db.get(Dataset, dataset_id)
    if d is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "数据集不存在")
    if not req.case_nos:
        return ok([])
    items = (
        db.query(Case)
        .filter(Case.dataset_id == dataset_id, Case.case_no.in_(req.case_nos))
        .order_by(Case.sort_order)
        .all()
    )
    return ok([CaseOut.model_validate(c).model_dump() for c in items])


@router.get("/{dataset_id}/download")
def download_dataset(dataset_id: uuid.UUID, db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    """导出数据集为 .xlsx：从 cases 按 dataset_id 检索，表头参照 column_mapping（标准字段→原表头）。"""
    d = db.get(Dataset, dataset_id)
    if d is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "数据集不存在")
    # column_mapping: {标准字段: 原表头}；表头按标准字段顺序取原表头名，缺省回退标准字段名
    mapping = d.column_mapping or {}
    fields = ["case_no", "user_input", "expected_gt", "dimension_l1", "dimension_l2"]
    headers = [mapping.get(f) or f for f in fields]

    cases = db.query(Case).filter(Case.dataset_id == dataset_id).order_by(Case.sort_order).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "数据集"
    ws.append(headers)
    for c in cases:
        ws.append([
            c.case_no or "", c.user_input or "", c.expected_gt or "",
            c.dimension_l1 or "", c.dimension_l2 or "",
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=dataset_{dataset_id}.xlsx"},
    )


@router.post("/upload", status_code=201)
async def upload(
    file: UploadFile = File(...),
    name: str = Form(...),
    column_mapping: str = Form(...),  # JSON 字符串
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    try:
        mapping = ColumnMapping(**json.loads(column_mapping))
    except Exception as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"列映射无效: {e}")
    # 校验必填映射
    for field in REQUIRED:
        if not getattr(mapping, field):
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"必填列未映射: {field}")

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "仅支持 .xlsx")
    wb = load_workbook(io.BytesIO(await file.read()), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "无数据行")
    header = [str(h) if h is not None else "" for h in rows[0]]
    idx = {name_: header.index(name_) for name_ in [mapping.case_no, mapping.user_input, mapping.expected_gt,
                                                     mapping.dimension_l1, mapping.dimension_l2] if name_}

    dataset = Dataset(name=name, owner_id=user.id, source="excel",
                      column_mapping=mapping.model_dump())
    db.add(dataset)
    db.flush()
    count = 0
    for i, row in enumerate(rows[1:], start=1):
        if row is None or all(v is None for v in row):
            continue
        def _get(field):
            col = getattr(mapping, field)
            if not col or col not in idx:
                return None
            v = row[idx[col]]
            return str(v) if v is not None else None
        case_no, user_input, expected_gt = _get("case_no"), _get("user_input"), _get("expected_gt")
        if not case_no or not user_input or not expected_gt:
            continue  # 跳过缺必填的行
        db.add(Case(
            dataset_id=dataset.id, case_no=case_no, user_input=user_input, expected_gt=expected_gt,
            dimension_l1=_get("dimension_l1"), dimension_l2=_get("dimension_l2"), sort_order=i,
        ))
        count += 1
    dataset.case_count = count
    db.commit()
    db.refresh(dataset)
    return ok(DatasetOut.model_validate(dataset).model_dump())


@router.post("/{dataset_id}/rename")
def rename_dataset(dataset_id: uuid.UUID, req: DatasetNameUpdateReq, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    d = db.get(Dataset, dataset_id)
    if d is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "数据集不存在")
    require_owner_or_admin(d.owner_id, user)  # 普通用户仅自己的，管理员全部
    d.name = req.name
    db.commit()
    db.refresh(d)
    return ok(DatasetOut.model_validate(d).model_dump())


def _delete_dataset_cascade(db: Session, dataset_id: uuid.UUID) -> None:
    """删数据集及其关联的 cases / case_results（先删 case_results，再删 cases，最后删 dataset）。"""
    case_ids = [c.id for c in db.query(Case).filter(Case.dataset_id == dataset_id)]
    if case_ids:
        run_ids = [r.run_id for r in db.query(CaseResult.run_id).filter(CaseResult.case_id.in_(case_ids)).distinct()]
        if run_ids:
            db.query(CaseResult).filter(CaseResult.run_id.in_(run_ids)).delete(synchronize_session=False)
        db.query(Case).filter(Case.id.in_(case_ids)).delete(synchronize_session=False)
    d = db.get(Dataset, dataset_id)
    if d is not None:
        db.delete(d)


@router.post("/{dataset_id}/delete")
def delete(dataset_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    d = db.get(Dataset, dataset_id)
    if d is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "数据集不存在")
    require_owner_or_admin(d.owner_id, user)  # 普通用户仅自己的，管理员全部
    _delete_dataset_cascade(db, dataset_id)
    db.commit()
    return ok()


@router.post("/batch-delete")
def batch_delete(req: DatasetBatchDeleteReq, db: Session = Depends(get_db), _admin: User = Depends(require_admin)):
    """管理员批量删除数据集，请求体 `{ids: [...]}`。"""
    if not req.ids:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "ids 不能为空")
    deleted = 0
    for raw in req.ids:
        try:
            dsid = uuid.UUID(raw)
        except ValueError:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"非法的数据集 id: {raw}")
        if db.get(Dataset, dsid) is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, f"数据集不存在: {raw}")
        _delete_dataset_cascade(db, dsid)
        deleted += 1
    db.commit()
    return ok({"deleted": deleted})
