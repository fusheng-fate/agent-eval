"""任务：创建（单用户单活跃任务 P25）/ 列表 / 详情 / 逐用例 / 暂停继续停止 / 导出 / 自检 / 看板。"""
import io
import uuid
from datetime import datetime, timezone

import json as _json
import time as _time

from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import Response, StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import func, insert, select, update
from sqlalchemy.orm import Session

from ...core.config import settings
from ...core.database import get_db
from ...core.deps import get_current_user, require_admin, require_owner_or_admin
from ...models import Case, CaseResult, Config, Dataset, Evaluator, FlowTemplate, Metric, Report, Run, RunLog, User
from ..shared import ApiResp, Msg, PageData, ok
from .export_service import export_run_excel
from .schemas import CaseResultOut, DashboardSummary, EvalImportColumnMapping, RunCreateReq, RunOut

router = APIRouter(prefix="/runs", tags=["runs"])
dashboard_router = APIRouter(prefix="/dashboard", tags=["dashboard"])

ACTIVE = ["pending", "running", "paused"]


def _utcnow():
    return datetime.now(timezone.utc)


def _natural_key(s: str):
    """自然排序 key：数字段转 int，非数字段原样。使 A01<A02<A10<B01。"""
    import re
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s) if t != ""]


def _get_config(db: Session, key: str, default):
    row = db.query(Config).filter(Config.config_key == key).first()
    if row and row.config_value:
        cv = row.config_value
        # 兼容两种存储格式：{"value": X}（seed/字符串配置）与裸 dict（前端保存的 llm.auth）
        if isinstance(cv, dict) and "value" in cv:
            return cv["value"]
        return cv
    return default


@router.post("", response_model=ApiResp[RunOut], status_code=201)
def create_run(req: RunCreateReq, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    # P25：同一用户只能有一个活跃任务
    active = db.query(Run).filter(Run.owner_id == user.id, Run.status.in_(ACTIVE)).first()
    if active:
        raise HTTPException(status.HTTP_409_CONFLICT, "已有任务进行中，请先完成或停止后再创建")

    dataset = db.get(Dataset, uuid.UUID(req.dataset_id))
    if dataset is None or dataset.case_count == 0:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "数据集不存在或为空")
    template = db.get(FlowTemplate, uuid.UUID(req.flow_template_id))
    if template is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "流程模板不存在")
    evaluator = db.query(Evaluator).filter(Evaluator.is_standard.is_(True)).first()
    if evaluator is None:
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, "标准评估器未初始化")

    # Q3：额外指标数量上限（管理员可在配置中心调整，默认 5）
    extra_max = int(_get_config(db, "evaluator.extra_metric_limit", 5))
    if len(req.extra_metric_ids) > extra_max:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            f"额外指标最多 {extra_max} 个（当前 {len(req.extra_metric_ids)}）",
        )

    # Q4：额外指标须为指标库中存在、且标准评估器未引用的指标
    std_metric_ids = {em.metric_id for em in evaluator.metrics}
    for mid in req.extra_metric_ids:
        metric = db.get(Metric, uuid.UUID(mid))
        if metric is None:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"额外指标不存在: {mid}")
        if metric.id in std_metric_ids:
            raise HTTPException(
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                f"额外指标不能是标准评估器已引用的指标: {metric.name}",
            )

    # P13：overrides 仅允许模板开放（user_editable=True）的参数；管理员可覆盖任意开放项。
    # 普通用户若传了未开放的 paramPath → 422。
    perms = {p.param_path: p.user_editable for p in template.param_perms}
    clean_overrides: dict = {}
    for path, value in (req.overrides or {}).items():
        if user.role == "admin":
            clean_overrides[path] = value
        elif perms.get(path, False):
            clean_overrides[path] = value
        else:
            raise HTTPException(
                status.HTTP_422_UNPROCESSABLE_ENTITY,
                f"参数 {path} 未开放给普通用户编辑",
            )

    # 轮次：round_size=1 等同于无轮次（每组 1 个 case，无串行依赖，无前置 API）
    round_size = req.round_size
    if round_size is not None and round_size < 1:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "round_size 必须为正整数")
    if round_size == 1:
        round_size = None  # 退化为无轮次模式
    total_rounds = None
    if round_size is not None:
        total_rounds = -(-dataset.case_count // round_size)  # ceil division

    run = Run(
        task_name=req.task_name, owner_id=user.id, evaluator_id=evaluator.id,
        dataset_id=dataset.id, flow_template_id=template.id,
        extra_metric_ids=[uuid.UUID(m) for m in req.extra_metric_ids],
        overrides=clean_overrides or None,
        status="pending", total_cases=dataset.case_count,
        model_concurrency=int(_get_config(db, "concurrency.model", settings.CONCURRENCY_MODEL)),
        target_concurrency=int(template.target_concurrency),
        round_size=round_size, total_rounds=total_rounds, done_rounds=0,
    )
    db.add(run)
    db.flush()

    # 为每条 case 批量建 case_result(pending)——PG 队列：pending 行即队列，
    # worker 直接抢，无需逐条入队 Redis（避免几万条 lpush 阻塞请求）
    cases = db.query(Case).filter(Case.dataset_id == dataset.id).order_by(Case.sort_order).all()
    if cases:
        # 计算 round_no（按 sort_order 顺序 + round_size 切分）
        for i, c in enumerate(cases):
            c.round_no = (i // round_size) if round_size else None
        db.flush()
        case_ids = [c.id for c in cases]
        db.execute(
            insert(CaseResult).values(
                [{"run_id": run.id, "case_id": cid, "status": "pending"} for cid in case_ids]
            )
        )
    run.status = "running"
    run.started_at = _utcnow()
    db.commit()
    db.refresh(run)
    return ok(RunOut.model_validate(run))


@router.post("/import", response_model=ApiResp[RunOut], status_code=201)
async def import_eval_results(
    file: UploadFile = File(...),
    task_name: str = Form(...),
    column_mapping: str = Form(...),  # JSON 字符串
    extra_metric_ids: str = Form("[]"),  # JSON 数组字符串
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """导入评测结果 Excel → 建任务异步打分（mode=eval_import，跳过调被测 Agent）。

    agent_output（执行结果列）解析后直接填进 case_results.agent_output，
    worker 对该类任务跳过 call_target_multi_turn，走原有 LLM 打分逻辑。
    后续查看结果/报告/导出/暂停续跑/看板与 exec 任务完全一致。
    """
    try:
        mapping = EvalImportColumnMapping(**_json.loads(column_mapping))
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"列映射无效: {e}")
    for field in ("case_no", "user_input", "expected_gt", "agent_output"):
        if not getattr(mapping, field):
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"必填列未映射: {field}")
    try:
        extra_ids = _json.loads(extra_metric_ids) if extra_metric_ids else []
    except _json.JSONDecodeError:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "extra_metric_ids 非合法 JSON 数组")

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "仅支持 .xlsx")

    # P25：同一用户只能有一个活跃任务（与 create_run 一致）
    active = db.query(Run).filter(Run.owner_id == user.id, Run.status.in_(ACTIVE)).first()
    if active:
        raise HTTPException(status.HTTP_409_CONFLICT, "已有任务进行中，请先完成或停止后再创建")

    evaluator = db.query(Evaluator).filter(Evaluator.is_standard.is_(True)).first()
    if evaluator is None:
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, "标准评估器未初始化")

    # Q3/Q4：额外指标上限 + 存在性 + 不与标准指标重复（与 create_run 一致）
    extra_max = int(_get_config(db, "evaluator.extra_metric_limit", 5))
    if len(extra_ids) > extra_max:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"额外指标最多 {extra_max} 个（当前 {len(extra_ids)}）")
    std_metric_ids = {em.metric_id for em in evaluator.metrics}
    for mid in extra_ids:
        metric = db.get(Metric, uuid.UUID(mid))
        if metric is None:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"额外指标不存在: {mid}")
        if metric.id in std_metric_ids:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"额外指标不能是标准评估器已引用的指标: {metric.name}")

    # 解析 Excel（复用 datasets 上传的列映射解析模式）
    wb = load_workbook(io.BytesIO(await file.read()), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "无数据行")
    header = [str(h) if h is not None else "" for h in rows[0]]
    fields = [mapping.case_no, mapping.user_input, mapping.expected_gt,
              mapping.agent_output, mapping.dimension_l1, mapping.dimension_l2]
    idx = {name_: header.index(name_) for name_ in fields if name_}

    def _get(row, field):
        col = getattr(mapping, field)
        if not col or col not in idx:
            return None
        v = row[idx[col]]
        return str(v) if v is not None else None

    # 建 Dataset + Cases
    dataset = Dataset(name=task_name, owner_id=user.id, source="eval_import",
                      column_mapping=mapping.model_dump())
    db.add(dataset)
    db.flush()

    parsed: list[dict] = []  # {case_id, agent_output}
    count = 0
    for i, row in enumerate(rows[1:], start=1):
        if row is None or all(v is None for v in row):
            continue
        case_no = _get(row, "case_no")
        user_input = _get(row, "user_input")
        expected_gt = _get(row, "expected_gt")
        agent_output = _get(row, "agent_output")
        if not case_no or not user_input or not expected_gt or not agent_output:
            continue  # 跳过缺必填的行
        case = Case(
            dataset_id=dataset.id, case_no=case_no, user_input=user_input,
            expected_gt=expected_gt,
            dimension_l1=_get(row, "dimension_l1"), dimension_l2=_get(row, "dimension_l2"),
            sort_order=i,
        )
        db.add(case)
        db.flush()
        parsed.append({"case_id": case.id, "agent_output": agent_output})
        count += 1

    if count == 0:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "无有效数据行（必填列齐全的行）")
    dataset.case_count = count
    db.flush()

    # 建 Run(mode=eval_import) + CaseResult(pending, agent_output 已填)
    run = Run(
        task_name=task_name, owner_id=user.id, evaluator_id=evaluator.id,
        dataset_id=dataset.id, flow_template_id=None, mode="eval_import",
        extra_metric_ids=[uuid.UUID(m) for m in extra_ids],
        status="pending", total_cases=count,
        model_concurrency=int(_get_config(db, "concurrency.model", settings.CONCURRENCY_MODEL)),
        target_concurrency=settings.CONCURRENCY_TARGET,
    )
    db.add(run)
    db.flush()
    if parsed:
        db.execute(
            insert(CaseResult).values(
                [{"run_id": run.id, "case_id": p["case_id"], "status": "pending",
                  "agent_output": p["agent_output"]} for p in parsed]
            )
        )
    run.status = "running"
    run.started_at = _utcnow()
    db.commit()
    db.refresh(run)
    return ok(RunOut.model_validate(run))


@router.get("", response_model=ApiResp[PageData[RunOut]])
def list_runs(
    status_: str | None = Query(None, alias="status"),
    owner: str | None = None,
    keyword: str | None = None,
    username: str | None = None,
    created_start: str | None = None,
    created_end: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    q = db.query(Run)
    if status_:
        q = q.filter(Run.status == status_)
    if owner == "mine":
        q = q.filter(Run.owner_id == _user.id)
    elif owner and owner != "all":
        q = q.filter(Run.owner_id == uuid.UUID(owner))
    if keyword:
        q = q.filter(Run.task_name.ilike(f"%{keyword}%"))
    if username:
        q = q.join(User, Run.owner_id == User.id).filter(User.username == username)
    if created_start:
        q = q.filter(Run.created_at >= created_start)
    if created_end:
        q = q.filter(Run.created_at <= created_end)
    total = q.count()
    runs = q.order_by(Run.created_at.desc()).offset((page - 1) * page_size).limit(page_size).all()
    # 批量查 owner 用户名 + 关联报告 ID
    owner_ids = {r.owner_id for r in runs}
    user_map = {u.id: u.username for u in db.query(User).filter(User.id.in_(owner_ids)).all()} if owner_ids else {}
    run_ids = {r.id for r in runs}
    report_map = {rep.run_id: rep.id for rep in db.query(Report.run_id, Report.id).filter(Report.run_id.in_(run_ids)).all()} if run_ids else {}
    items = []
    for r in runs:
        out = RunOut.model_validate(r)
        out.username = user_map.get(r.owner_id)
        out.report_id = str(report_map[r.id]) if r.id in report_map else None
        items.append(out)
    return ok(PageData(items=items, page=page, page_size=page_size, total=total))


@router.get("/{run_id}", response_model=ApiResp[RunOut])
def get_run(run_id: uuid.UUID, db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")
    out = RunOut.model_validate(run)
    rep = db.query(Report.id).filter(Report.run_id == run_id).first()
    out.report_id = str(rep[0]) if rep else None
    return ok(out)


@router.get("/{run_id}/cases", response_model=ApiResp[PageData[CaseResultOut]])
def list_case_results(run_id: uuid.UUID, page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100), db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")
    # 按用例名自然排序（A01<A02<A10<B01），case_no 在 Case 表，需 join 后内存排序
    rows = (
        db.query(CaseResult, Case.case_no)
        .join(Case, CaseResult.case_id == Case.id)
        .filter(CaseResult.run_id == run_id)
        .all()
    )
    total = len(rows)
    rows.sort(key=lambda r: _natural_key(r[1] or ""))
    results = [r[0] for r in rows[(page - 1) * page_size : page * page_size]]
    case_map = {c.id: c for c in db.query(Case).filter(Case.id.in_([r.case_id for r in results])).all()}
    out = []
    for r in results:
        c = case_map.get(r.case_id)
        d = CaseResultOut.model_validate(r)
        d.case_no = c.case_no if c else None
        out.append(d)
    return ok(PageData(items=out, page=page, page_size=page_size, total=total))


@router.get("/{run_id}/cases/{case_id}", response_model=ApiResp[CaseResultOut])
def get_case_result(run_id: uuid.UUID, case_id: uuid.UUID, db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")
    cr = db.query(CaseResult).filter(CaseResult.run_id == run_id, CaseResult.case_id == case_id).first()
    if cr is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "用例结果不存在")
    c = db.get(Case, case_id)
    d = CaseResultOut.model_validate(cr)
    d.case_no = c.case_no if c else None
    return ok(d)


@router.get("/{run_id}/logs")
def list_run_logs(run_id: uuid.UUID, db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    """任务执行日志（terminal 风格展示用），按时间正序。"""
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")
    logs = (
        db.query(RunLog)
        .filter(RunLog.run_id == run_id)
        .order_by(RunLog.created_at.asc())
        .all()
    )
    return ok([
        {"id": str(l.id), "level": l.level, "message": l.message, "created_at": l.created_at.isoformat()}
        for l in logs
    ])


def _set_status(run: Run, new_status: str, db: Session) -> None:
    if new_status == "paused" and run.status != "running":
        raise HTTPException(status.HTTP_409_CONFLICT, "仅执行中任务可暂停")
    if new_status == "running" and run.status not in ("paused",):
        raise HTTPException(status.HTTP_409_CONFLICT, "仅已暂停任务可继续")
    run.status = new_status
    if new_status in ("stopped", "error") and run.finished_at is None:
        run.finished_at = _utcnow()
    db.commit()


@router.post("/{run_id}/pause", response_model=ApiResp[RunOut])
def pause(run_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")
    require_owner_or_admin(run.owner_id, user)
    _set_status(run, "paused", db)
    return ok(RunOut.model_validate(run))


@router.post("/{run_id}/resume", response_model=ApiResp[RunOut])
def resume(run_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")
    require_owner_or_admin(run.owner_id, user)
    # 无需重置在途 case：pause 后未开始的 case 已由 worker 入口自行回退，
    # 在途执行/评分的 worker 会自然跑完；真正卡死的（worker 崩溃）由
    # recover_stale_running（心跳意识）回收。这里只把 run 恢复为 running。
    # 若在此无条件 running→pending / scoring→executed，会与在途 worker 竞态，
    # 造成同一 case 被重复执行/重复打分。
    _set_status(run, "running", db)
    return ok(RunOut.model_validate(run))


@router.post("/{run_id}/stop", response_model=ApiResp[RunOut])
def stop(run_id: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")
    require_owner_or_admin(run.owner_id, user)
    # 停止是终态：把该 run 下所有未完成 case 批量置 skipped（清锁），
    # 否则 worker 回退的 case 会永久卡在中间态。
    # 中间态含执行/评分全链路：pending/running/executed/scoring。
    db.execute(
        update(CaseResult)
        .where(CaseResult.run_id == run_id, CaseResult.status.in_(["pending", "running", "executed", "scoring"]))
        .values(status="skipped", locked_by=None, locked_at=None)
    )
    # 同步计数（skipped 是终态，计入 done）
    from .worker import _sync_run_counters
    _sync_run_counters(db, run_id)
    _set_status(run, "stopped", db)
    return ok(RunOut.model_validate(run))


@router.post("/{run_id}/delete", response_model=ApiResp[None])
def delete_run(run_id: uuid.UUID, db: Session = Depends(get_db), _admin: User = Depends(require_admin)):
    # Q2 已定：仅管理员可删除任务，普通用户仅可暂停/继续/停止
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")
    if run.status in ACTIVE:
        raise HTTPException(status.HTTP_409_CONFLICT, "任务进行中，请先停止")
    db.delete(run)
    db.commit()
    return ok(None)


@router.get("/{run_id}/export")
def export(run_id: uuid.UUID, level: str = Query("eval", pattern="^(exec|eval)$"), db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")
    data = export_run_excel(db, run, level)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=run_{run_id}.xlsx"},
    )


@router.get("/{run_id}/events")
async def events(run_id: uuid.UUID, db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    """SSE 实时事件流：case 完成 / 任务完成时推送。

    前端用 EventSource 订阅，收到 case_done 事件刷新进度，run_done 事件停止。
    """
    run = db.get(Run, run_id)
    if run is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "任务不存在")

    def _gen():
        from ...core.database import SessionLocal
        from ...core.redis_client import get_redis
        pubsub = get_redis().pubsub()
        pubsub.subscribe(f"agent_eval:events:{run_id}")
        sdb = SessionLocal()
        try:
            # 先推送当前状态（连接即得）
            cur = sdb.execute(
                select(Run.status, Run.done_cases, Run.total_cases, Run.passed_cases, Run.failed_cases)
                .where(Run.id == run_id)
            ).one()
            yield f"data: {_json.dumps({'type': 'snapshot', 'status': cur[0], 'done': cur[1], 'total': cur[2], 'passed': cur[3], 'failed': cur[4]})}\n\n"
            if cur[0] in ("done", "stopped", "error"):
                return
            while True:
                msg = pubsub.get_message(timeout=15)
                if msg and msg["type"] == "message":
                    yield f"data: {msg['data']}\n\n"
                    try:
                        evt = _json.loads(msg["data"])
                        if evt.get("type") == "run_done":
                            break
                    except Exception:
                        pass
                else:
                    yield ": heartbeat\n\n"
        finally:
            pubsub.unsubscribe(f"agent_eval:events:{run_id}")
            pubsub.close()
            sdb.close()

    return StreamingResponse(_gen(), media_type="text/event-stream", headers={
        "Cache-Control": "no-cache",
        "Connection": "keep-alive",
        "X-Accel-Buffering": "no",
    })


@router.post("/selfcheck/exec", response_model=ApiResp[dict])
def selfcheck_exec(req: RunCreateReq, db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    """执行自检：取第一个用例真实调被测 Agent，验证流程模板链路连通。"""
    template = db.get(FlowTemplate, uuid.UUID(req.flow_template_id))
    if template is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "流程模板不存在")
    dataset = db.get(Dataset, uuid.UUID(req.dataset_id))
    if dataset is None or dataset.case_count == 0:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "数据集为空")
    first = (
        db.query(Case).filter(Case.dataset_id == dataset.id).order_by(Case.sort_order).first()
    )
    if first is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "数据集为空")
    row = {
        "case_id": str(first.id), "case_no": first.case_no, "user_input": first.user_input,
        "expected_gt": first.expected_gt,
        "dimension_l1": first.dimension_l1 or "", "dimension_l2": first.dimension_l2 or "",
    }
    from .target_client import (
        _extract_path, apply_overrides, call_pre_api, call_target_multi_turn, pre_api_enabled,
    )
    trace: list = []
    chain = apply_overrides(template.chain_json, req.overrides)
    # 前置 API（v1.5+）：自检也须真实触发，否则主 API 引用的 {{session_id}} 等占位符
    # 渲染不出，自检给出假阳性。取变量组第一组、不写 Redis 缓存（一次性连通验证）。
    if pre_api_enabled(chain):
        pre_api = chain.get("pre_api")
        groups = chain.get("pre_api_groups") or []
        group_vars = dict(groups[0]) if groups else {}
        pre_out = call_pre_api(
            pre_api,
            variables={**(chain.get("variables") or {}), **group_vars},
            trace=trace,
        )
        for var_name, json_path in (pre_api.get("extract") or {}).items():
            row[var_name] = _extract_path(pre_out, json_path)
        row.update(group_vars)
    try:
        out = call_target_multi_turn(chain, row, trace=trace)
    except Exception as e:  # noqa: BLE001
        # 失败时把已记录的 trace 一并返回，便于定位是哪一步、什么请求/响应出的问题
        return ok({"message": f"被测 Agent 调用失败: {e}", "ok": False, "trace": trace})
    preview = str(out)[:200]
    return ok({"message": f"执行自检通过：被测 Agent 已连通，返回预览 {preview!r}", "ok": True, "trace": trace})


@router.post("/selfcheck/eval", response_model=ApiResp[dict])
def selfcheck_eval(db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    """评测自检：测 LLM 接口是否正常。"""
    from .llm_client import call_llm, resolve_llm_token
    cfg = {
        "base_url": _get_config(db, "llm.base_url", settings.LLM_BASE_URL),
        "api_key": _get_config(db, "llm.api_key", settings.LLM_API_KEY),
        "model": _get_config(db, "llm.model", settings.LLM_MODEL),
        "timeout": 30,
        "auth": _get_config(db, "llm.auth", None),
    }
    if not cfg["base_url"]:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "LLM base_url 未配置")
    try:
        if cfg["auth"] is not None:
            cfg["_token"], cfg["_token_query"] = resolve_llm_token(db, cfg)
            cfg["_db"] = db
        call_llm(cfg, "回复 OK", parse=False)
        return ok({"message": "评测自检通过：LLM 接口正常"})
    except Exception as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"LLM 接口异常: {e}")


@router.post("/selfcheck/llm-token", response_model=ApiResp[dict])
def selfcheck_llm_token(
    auth_override: dict | None = Body(None),
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    """手动刷新动态 token：强制绕过缓存调 token 接口，验证取 token 链路是否生效。

    仅 dynamic 认证可用；静态/未配置返回提示。成功返回脱敏后的 token 预览 + 缓存 TTL。

    auth_override: 可选（请求体），前端表单未保存的草稿（AuthConfig）。传入时以草稿为准测试，
    让用户在保存前即可验证配置是否正确；不传则用已保存的 llm.auth。
    """
    from .llm_client import invalidate_llm_token_cache, resolve_llm_token
    # 优先用前端草稿（未保存也能验证），否则回退到已保存配置
    # Body 无 body 时可能传 {} 或 None，统一视为"未传草稿"
    auth = (auth_override if auth_override else None) or _get_config(db, "llm.auth", None)
    if not auth or auth.get("type") != "dynamic":
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "当前非动态认证，无需刷新 token")
    cfg = {
        "api_key": _get_config(db, "llm.api_key", settings.LLM_API_KEY),
        "auth": auth,
    }
    try:
        invalidate_llm_token_cache(auth)  # 强制重取，验证真实接口
        token, _ = resolve_llm_token(db, cfg)
    except Exception as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"动态 token 获取失败: {e}")
    masked = (token[:6] + "…" + token[-4:]) if len(token) > 12 else "****"
    ta = auth.get("token_api") or {}
    ttl = int(ta.get("ttl") or 300)
    return ok({
        "message": "动态 token 获取成功，已写回 API 密钥并缓存",
        "token_preview": masked,
        "cache_ttl": max(1, ttl - 60),
    })


# ---------- 看板（P3，任务中心顶部） ----------
@dashboard_router.get("/summary", response_model=ApiResp[DashboardSummary])
def summary(
    since: datetime | None = None,
    until: datetime | None = None,
    owner: str | None = Query(None, description="按用户筛选：mine=仅自己，UUID=指定用户，缺省=全局"),
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    # Q5：支持按时间范围 / 按用户筛选，默认全局
    def _owner_filter(q):
        if owner == "mine":
            return q.filter(Run.owner_id == _user.id)
        if owner and owner != "all":
            return q.filter(Run.owner_id == uuid.UUID(owner))
        return q

    q = _owner_filter(db.query(CaseResult).join(Run, CaseResult.run_id == Run.id))
    if since:
        q = q.filter(Run.created_at >= since)
    if until:
        q = q.filter(Run.created_at <= until)

    total = q.count()
    passed = q.filter(CaseResult.status == "passed").count()
    failed = q.filter(CaseResult.status == "failed").count()
    tokens = _owner_filter(
        db.query(func.coalesce(func.sum(CaseResult.token_usage), 0)).join(Run, CaseResult.run_id == Run.id)
    ).scalar()
    lat_q = _owner_filter(
        db.query(func.avg(CaseResult.latency_sec)).join(Run, CaseResult.run_id == Run.id)
    )
    lat = lat_q.filter(CaseResult.latency_sec.isnot(None)).scalar()

    run_q = _owner_filter(db.query(Run))
    if since:
        run_q = run_q.filter(Run.created_at >= since)
    if until:
        run_q = run_q.filter(Run.created_at <= until)
    total_runs = run_q.count()
    running_runs = run_q.filter(Run.status.in_(["pending", "running", "paused"])).count()

    return ok(DashboardSummary(
        total_cases=total, passed_cases=passed, failed_cases=failed,
        pass_rate=round(passed / total, 4) if total else 0,
        total_runs=total_runs, running_runs=running_runs,
        token_usage=int(tokens or 0), avg_latency_sec=round(float(lat), 2) if lat else 0,
    ))
